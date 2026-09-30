import io
from sqlalchemy.orm import Session
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

from app.models.models import AuditJob, Finding

def generate_pdf_report(job: AuditJob) -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36
    )

    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=20,
        leading=24,
        textColor=colors.HexColor('#0f172a'),
        spaceAfter=6
    )

    subtitle_style = ParagraphStyle(
        'DocSubtitle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        textColor=colors.HexColor('#64748b'),
        spaceAfter=15
    )

    section_style = ParagraphStyle(
        'SectionHeader',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=13,
        leading=16,
        textColor=colors.HexColor('#1e293b'),
        spaceBefore=12,
        spaceAfter=8
    )

    cell_style = ParagraphStyle(
        'TableCell',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=11,
        textColor=colors.HexColor('#334155')
    )

    cell_bold = ParagraphStyle(
        'TableCellBold',
        parent=cell_style,
        fontName='Helvetica-Bold',
        textColor=colors.HexColor('#0f172a')
    )

    elements = []

    # Title Banner
    elements.append(Paragraph("AI-Driven Network Security Compliance Audit Report", title_style))
    elements.append(Paragraph(f"NTRO SIH26155 • Job ID: {job.job_number} • Vendor: {job.vendor.upper()} • Host: {job.hostname}", subtitle_style))
    elements.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor('#00f2fe'), spaceAfter=15))

    # Executive Summary Card
    summary_data = [
        [Paragraph("Compliance Score", cell_bold), Paragraph(f"{job.compliance_score:.1f}%", cell_bold)],
        [Paragraph("Audit Status", cell_style), Paragraph(job.status, cell_style)],
        [Paragraph("Target Hostname", cell_style), Paragraph(job.hostname, cell_style)],
        [Paragraph("Total Security Controls Checked", cell_style), Paragraph(str(job.total_rules), cell_style)],
        [Paragraph("Passed Controls", cell_style), Paragraph(f"{job.passed_rules} (PASS)", cell_style)],
        [Paragraph("Failed / Non-Compliant Controls", cell_style), Paragraph(f"{job.failed_rules} (FAIL)", cell_style)],
        [Paragraph("Audit Date", cell_style), Paragraph(job.created_at.strftime("%Y-%m-%d %H:%M UTC"), cell_style)],
    ]

    t_summary = Table(summary_data, colWidths=[200, 320])
    t_summary.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#f8fafc')),
        ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#cbd5e1')),
        ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor('#e2e8f0')),
        ('PADDING', (0,0), (-1,-1), 6),
    ]))
    elements.append(t_summary)
    elements.append(Spacer(1, 15))

    # Detailed Findings Section
    elements.append(Paragraph("Detailed Security Findings & Framework Mapping", section_style))

    findings_table_data = [
        [
            Paragraph("Rule & Title", cell_bold),
            Paragraph("Severity", cell_bold),
            Paragraph("Status", cell_bold),
            Paragraph("Evidence Line", cell_bold),
            Paragraph("Framework Mappings", cell_bold)
        ]
    ]

    for f in job.findings:
        sev_color = "#ef4444" if f.severity == "CRITICAL" else "#f59e0b" if f.severity == "HIGH" else "#3b82f6"
        stat_color = "#10b981" if f.status == "PASS" else "#ef4444"

        rule_p = Paragraph(f"<b>{f.rule_code}</b><br/>{f.title}", cell_style)
        sev_p = Paragraph(f"<font color='{sev_color}'><b>{f.severity}</b></font>", cell_style)
        stat_p = Paragraph(f"<font color='{stat_color}'><b>{f.status}</b></font>", cell_style)
        ev_p = Paragraph(f"Line {f.evidence_start_line}: <i>{f.evidence_raw or 'N/A'}</i>", cell_style)
        
        fw_text = f"CIS: {f.cis_mapping or 'N/A'}<br/>NIST: {f.nist_mapping or 'N/A'}<br/>STIG: {f.stig_mapping or 'N/A'}"
        fw_p = Paragraph(fw_text, cell_style)

        findings_table_data.append([rule_p, sev_p, stat_p, ev_p, fw_p])

    t_findings = Table(findings_table_data, colWidths=[130, 60, 50, 140, 140])
    t_findings.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#0f172a')),
        ('TEXTCOLOR', (0,0), (-1,0), colors.white),
        ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#cbd5e1')),
        ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor('#e2e8f0')),
        ('PADDING', (0,0), (-1,-1), 5),
        ('VALIGN', (0,0), (-1,-1), 'TOP'),
    ]))

    elements.append(t_findings)
    doc.build(elements)
    buffer.seek(0)
    return buffer.getvalue()


def generate_excel_report(job: AuditJob) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Audit Executive Summary"

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")

    ws.append(["Audit Job Number", job.job_number])
    ws.append(["Target Device", job.hostname])
    ws.append(["Vendor & OS", job.vendor.upper()])
    ws.append(["Compliance Score", f"{job.compliance_score:.1f}%"])
    ws.append(["Passed Rules", job.passed_rules])
    ws.append(["Failed Rules", job.failed_rules])
    ws.append(["Audit Timestamp", job.created_at.isoformat()])
    ws.append([])

    ws_findings = wb.create_sheet(title="Detailed Findings")
    headers = [
        "Rule Code", "Title", "Category", "Severity", "Status",
        "Start Line", "Evidence Raw", "Remediation", "CIS Benchmark", "NIST SP 800-53", "DISA STIG", "ISO 27001"
    ]
    ws_findings.append(headers)

    for col_num in range(1, len(headers) + 1):
        cell = ws_findings.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill

    for f in job.findings:
        ws_findings.append([
            f.rule_code,
            f.title,
            f.category,
            f.severity,
            f.status,
            f.evidence_start_line,
            f.evidence_raw or "",
            f.remediation or "",
            f.cis_mapping or "",
            f.nist_mapping or "",
            f.stig_mapping or "",
            f.iso_mapping or ""
        ])

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
